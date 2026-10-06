"""Create reproducible analysis deliverables for the synthetic MFG-005 data."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "mfg005_synthetic_line_performance.csv"
OUTPUT = ROOT / "outputs"
REDUCTIONS = (0.05, 0.10, 0.15)
DOWNTIME_FIELDS = {
    "downtime_category_mechanical": "Mechanical",
    "downtime_category_electrical": "Electrical",
    "downtime_category_tooling": "Tooling",
    "downtime_category_material": "Material",
    "downtime_category_operator": "Operator",
    "downtime_category_quality_hold": "Quality hold",
}
REQUIRED = {
    "shift_id",
    "line_id",
    "industry_sector",
    "shift_date",
    "takt_time_seconds",
    "actual_cycle_time_avg_seconds",
    "bottleneck_cycle_time_seconds",
    "bottleneck_station_id",
    "actual_production_quantity",
    "good_units_produced",
    "defective_units_produced",
    "throughput_rate_actual_uph",
    "oee_availability",
    "oee_performance",
    "oee_quality",
    "oee_overall",
    "oee_loss_primary_driver",
    "available_time_minutes",
    "downtime_unplanned_minutes",
    "downtime_planned_minutes",
    "downtime_event_count",
    "mtbf_minutes",
    "mttr_minutes",
    "wip_queue_avg_units",
    "wip_queue_max_units",
    "changeover_time_minutes",
    "changeover_count",
    "equipment_failure_flag",
    "maintenance_type_last",
    "equipment_age_years",
    "equipment_condition_score",
    "operator_utilisation_pct",
    "defect_rate_ppm",
}


def load_dataset() -> tuple[pd.DataFrame, pd.DataFrame, dict[str, object]]:
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Input dataset not found: {SOURCE}")
    source = pd.read_csv(SOURCE)
    missing = REQUIRED.difference(source.columns)
    if missing:
        raise ValueError(f"Input dataset is missing required columns: {sorted(missing)}")

    parsed = source.copy()
    parsed["shift_date"] = pd.to_datetime(
        parsed["shift_date"], format="%d-%m-%Y", errors="coerce"
    )
    auto = parsed[
        parsed["industry_sector"].astype("string").str.strip().str.casefold().eq("automotive")
    ].copy()
    if auto.empty:
        raise ValueError("No records match industry_sector='automotive'.")

    auto["bottleneck_gap_seconds"] = (
        auto["bottleneck_cycle_time_seconds"] - auto["takt_time_seconds"]
    )
    auto["actual_cycle_gap_seconds"] = (
        auto["actual_cycle_time_avg_seconds"] - auto["takt_time_seconds"]
    )
    auto["meets_takt"] = auto["bottleneck_gap_seconds"] <= 0
    auto["defect_rate_pct"] = (
        auto["defective_units_produced"]
        .div(auto["actual_production_quantity"].replace(0, pd.NA))
        .mul(100)
    )

    all_keys = source.groupby(["shift_id", "line_id"], dropna=False).size()
    auto_keys = auto.groupby(["shift_id", "line_id"], dropna=False).size()
    quality: dict[str, object] = {
        "source_rows": int(len(source)),
        "source_columns": int(source.shape[1]),
        "exact_duplicate_rows": int(source.duplicated().sum()),
        "duplicate_shift_line_keys": int(source.duplicated(["shift_id", "line_id"]).sum()),
        "duplicate_shift_line_key_groups": int((all_keys > 1).sum()),
        "automotive_duplicate_shift_line_records": int(
            auto.duplicated(["shift_id", "line_id"]).sum()
        ),
        "automotive_duplicate_shift_line_key_groups": int((auto_keys > 1).sum()),
        "automotive_unique_shift_line_keys": int(len(auto_keys)),
        "source_missing_cells": int(source.isna().sum().sum()),
        "source_missing_by_column": {
            str(column): int(count)
            for column, count in source.isna().sum().items()
            if count
        },
        "source_dtype_counts": {
            str(dtype): int(count)
            for dtype, count in source.dtypes.astype(str).value_counts().items()
        },
        "invalid_shift_dates": int(parsed["shift_date"].isna().sum()),
        "automotive_rows": int(len(auto)),
        "automotive_lines": int(auto["line_id"].nunique()),
        "industry_row_counts": {
            str(industry): int(count)
            for industry, count in source["industry_sector"].value_counts(dropna=False).items()
        },
    }
    return source, auto, quality


def line_summary(auto: pd.DataFrame) -> pd.DataFrame:
    result = auto.groupby("line_id", as_index=False).agg(
        observation_count=("shift_id", "count"),
        dominant_bottleneck_station=(
            "bottleneck_station_id",
            lambda s: str(s.dropna().mode().iloc[0]) if not s.dropna().mode().empty else "",
        ),
        mean_takt_seconds=("takt_time_seconds", "mean"),
        mean_actual_cycle_seconds=("actual_cycle_time_avg_seconds", "mean"),
        mean_bottleneck_cycle_seconds=("bottleneck_cycle_time_seconds", "mean"),
        mean_bottleneck_gap_seconds=("bottleneck_gap_seconds", "mean"),
        median_bottleneck_gap_seconds=("bottleneck_gap_seconds", "median"),
        records_over_takt_pct=("meets_takt", lambda s: (1 - s.mean()) * 100),
        mean_bottleneck_utilisation_pct=("bottleneck_utilisation_pct", "mean"),
        mean_oee=("oee_overall", "mean"),
        mean_availability=("oee_availability", "mean"),
        mean_performance=("oee_performance", "mean"),
        mean_quality=("oee_quality", "mean"),
        mean_throughput_uph=("throughput_rate_actual_uph", "mean"),
        mean_actual_units=("actual_production_quantity", "mean"),
        mean_good_units=("good_units_produced", "mean"),
        total_actual_units=("actual_production_quantity", "sum"),
        total_good_units=("good_units_produced", "sum"),
        total_defective_units=("defective_units_produced", "sum"),
        mean_wip_queue_units=("wip_queue_avg_units", "mean"),
        mean_max_wip_queue_units=("wip_queue_max_units", "mean"),
        total_unplanned_downtime_minutes=("downtime_unplanned_minutes", "sum"),
        total_planned_downtime_minutes=("downtime_planned_minutes", "sum"),
        total_downtime_events=("downtime_event_count", "sum"),
        mean_mtbf_minutes=("mtbf_minutes", "mean"),
        mean_mttr_minutes=("mttr_minutes", "mean"),
        mean_changeover_time_minutes=("changeover_time_minutes", "mean"),
        mean_changeover_count=("changeover_count", "mean"),
        mean_equipment_age_years=("equipment_age_years", "mean"),
        mean_equipment_condition_score=("equipment_condition_score", "mean"),
        mean_operator_utilisation_pct=("operator_utilisation_pct", "mean"),
    )
    positive = (
        auto.assign(positive_gap=auto["bottleneck_gap_seconds"].clip(lower=0))
        .groupby("line_id")["positive_gap"]
        .sum()
    )
    result["positive_gap_seconds_sum"] = result["line_id"].map(positive)
    result["bottleneck_severity_rank"] = (
        result["mean_bottleneck_gap_seconds"].rank(method="min", ascending=False).astype(int)
    )
    result = result.sort_values(
        ["mean_bottleneck_gap_seconds", "records_over_takt_pct"],
        ascending=[False, False],
    ).reset_index(drop=True)
    total_positive = result["positive_gap_seconds_sum"].sum()
    result["positive_gap_pareto_share_pct"] = (
        result["positive_gap_seconds_sum"] / total_positive * 100 if total_positive else 0
    )
    result["positive_gap_cumulative_pct"] = result["positive_gap_pareto_share_pct"].cumsum()
    return result


def station_summary(auto: pd.DataFrame) -> pd.DataFrame:
    result = auto.groupby("bottleneck_station_id", as_index=False).agg(
        bottleneck_observation_count=("shift_id", "count"),
        mean_bottleneck_cycle_seconds=("bottleneck_cycle_time_seconds", "mean"),
        mean_takt_seconds=("takt_time_seconds", "mean"),
        mean_bottleneck_gap_seconds=("bottleneck_gap_seconds", "mean"),
        mean_bottleneck_utilisation_pct=("bottleneck_utilisation_pct", "mean"),
        mean_oee=("oee_overall", "mean"),
    ).sort_values("bottleneck_observation_count", ascending=False, ignore_index=True)
    result["share_of_bottleneck_observations_pct"] = (
        result["bottleneck_observation_count"]
        / result["bottleneck_observation_count"].sum()
        * 100
    )
    result["cumulative_share_pct"] = result["share_of_bottleneck_observations_pct"].cumsum()
    return result


def downtime_pareto(auto: pd.DataFrame) -> pd.DataFrame:
    available = [column for column in DOWNTIME_FIELDS if column in auto.columns]
    if not available:
        raise ValueError("No downtime category fields were found.")
    result = (
        auto[available]
        .sum()
        .rename("downtime_minutes")
        .rename_axis("source_column")
        .reset_index()
    )
    result["category"] = result["source_column"].map(DOWNTIME_FIELDS)
    result = result.sort_values("downtime_minutes", ascending=False, ignore_index=True)
    category_total = result["downtime_minutes"].sum()
    result["share_of_categorized_downtime_pct"] = (
        result["downtime_minutes"] / category_total * 100 if category_total else 0
    )
    result["cumulative_share_pct"] = result["share_of_categorized_downtime_pct"].cumsum()
    result["total_unplanned_downtime_minutes"] = auto["downtime_unplanned_minutes"].sum()
    result["unplanned_minus_categorized_minutes"] = (
        result["total_unplanned_downtime_minutes"] - category_total
    )
    return result[
        [
            "category",
            "downtime_minutes",
            "share_of_categorized_downtime_pct",
            "cumulative_share_pct",
            "total_unplanned_downtime_minutes",
            "unplanned_minus_categorized_minutes",
        ]
    ]


def oee_tables(auto: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    labels = {
        "oee_availability": "Availability",
        "oee_performance": "Performance",
        "oee_quality": "Quality",
    }
    factors = pd.DataFrame(
        [
            {
                "factor": label,
                "mean_factor": auto[column].mean(),
                "mean_loss_to_100_pct": (1 - auto[column].mean()) * 100,
            }
            for column, label in labels.items()
        ]
    ).sort_values("mean_loss_to_100_pct", ascending=False, ignore_index=True)
    drivers = (
        auto["oee_loss_primary_driver"]
        .fillna("not recorded")
        .value_counts()
        .rename_axis("recorded_primary_loss_driver")
        .rename("observation_count")
        .reset_index()
    )
    drivers["share_of_observations_pct"] = drivers["observation_count"] / len(auto) * 100
    return factors, drivers


def simulation_table(auto: pd.DataFrame) -> pd.DataFrame:
    parts: list[pd.DataFrame] = []
    for reduction in REDUCTIONS:
        frame = auto.copy()
        frame["scenario_bottleneck_seconds"] = (
            frame["bottleneck_cycle_time_seconds"] * (1 - reduction)
        )
        frame["scenario_gap_seconds"] = (
            frame["scenario_bottleneck_seconds"] - frame["takt_time_seconds"]
        )
        frame["baseline_capacity_uph"] = 3600 / frame["bottleneck_cycle_time_seconds"]
        frame["scenario_capacity_uph"] = 3600 / frame["scenario_bottleneck_seconds"]
        grouped = frame.groupby("line_id", as_index=False).agg(
            observation_count=("shift_id", "count"),
            mean_takt_seconds=("takt_time_seconds", "mean"),
            mean_baseline_bottleneck_seconds=("bottleneck_cycle_time_seconds", "mean"),
            mean_scenario_bottleneck_seconds=("scenario_bottleneck_seconds", "mean"),
            mean_baseline_gap_seconds=("bottleneck_gap_seconds", "mean"),
            mean_scenario_gap_seconds=("scenario_gap_seconds", "mean"),
            records_meeting_scenario_takt_pct=(
                "scenario_gap_seconds",
                lambda s: (s <= 0).mean() * 100,
            ),
            baseline_theoretical_capacity_uph=("baseline_capacity_uph", "mean"),
            scenario_theoretical_capacity_uph=("scenario_capacity_uph", "mean"),
            observed_throughput_uph_mean=("throughput_rate_actual_uph", "mean"),
        )
        grouped["scenario_reduction_pct"] = reduction * 100
        grouped["hypothetical_capacity_delta_uph"] = (
            grouped["scenario_theoretical_capacity_uph"]
            - grouped["baseline_theoretical_capacity_uph"]
        )
        grouped["hypothetical_capacity_delta_pct"] = (
            grouped["hypothetical_capacity_delta_uph"]
            / grouped["baseline_theoretical_capacity_uph"]
            * 100
        )
        parts.append(grouped)
    result = pd.concat(parts, ignore_index=True)
    return result[
        [
            "line_id",
            "scenario_reduction_pct",
            "observation_count",
            "mean_takt_seconds",
            "mean_baseline_bottleneck_seconds",
            "mean_scenario_bottleneck_seconds",
            "mean_baseline_gap_seconds",
            "mean_scenario_gap_seconds",
            "records_meeting_scenario_takt_pct",
            "baseline_theoretical_capacity_uph",
            "scenario_theoretical_capacity_uph",
            "hypothetical_capacity_delta_uph",
            "hypothetical_capacity_delta_pct",
            "observed_throughput_uph_mean",
        ]
    ]


def exploratory_tables(auto: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    numeric = [
        "bottleneck_gap_seconds",
        "bottleneck_cycle_time_seconds",
        "downtime_unplanned_minutes",
        "equipment_age_years",
        "equipment_condition_score",
        "operator_utilisation_pct",
        "changeover_time_minutes",
        "changeover_count",
        "defect_rate_ppm",
        "oee_overall",
    ]
    correlations = auto[numeric].corr(method="spearman").reset_index(names="measure")
    rows: list[dict[str, object]] = []
    for column in ("equipment_failure_flag", "maintenance_type_last", "oee_loss_primary_driver"):
        grouped = auto.groupby(column, dropna=False).agg(
            observation_count=("shift_id", "count"),
            mean_bottleneck_gap_seconds=("bottleneck_gap_seconds", "mean"),
            mean_unplanned_downtime_minutes=("downtime_unplanned_minutes", "mean"),
            mean_oee=("oee_overall", "mean"),
            mean_equipment_age_years=("equipment_age_years", "mean"),
            mean_equipment_condition_score=("equipment_condition_score", "mean"),
        )
        for group, values in grouped.iterrows():
            rows.append(
                {
                    "grouping_variable": column,
                    "group": "not recorded" if pd.isna(group) else str(group),
                    **values.to_dict(),
                }
            )
    return correlations, pd.DataFrame(rows)


def daily_trend(auto: pd.DataFrame) -> pd.DataFrame:
    return (
        auto.groupby("shift_date", as_index=False)
        .agg(
            record_count=("shift_id", "count"),
            actual_units=("actual_production_quantity", "sum"),
            good_units=("good_units_produced", "sum"),
            defective_units=("defective_units_produced", "sum"),
            mean_oee=("oee_overall", "mean"),
            mean_takt_seconds=("takt_time_seconds", "mean"),
            mean_bottleneck_cycle_seconds=("bottleneck_cycle_time_seconds", "mean"),
            mean_throughput_uph=("throughput_rate_actual_uph", "mean"),
            unplanned_downtime_minutes=("downtime_unplanned_minutes", "sum"),
        )
        .sort_values("shift_date")
    )


def write_findings(
    auto: pd.DataFrame,
    quality: dict[str, object],
    lines: pd.DataFrame,
    down: pd.DataFrame,
    factors: pd.DataFrame,
    drivers: pd.DataFrame,
    scenarios: pd.DataFrame,
) -> None:
    top = lines.iloc[0]
    lead_down = down.iloc[0]
    lead_factor = factors.iloc[0]
    lead_driver = drivers.iloc[0]
    ten_pct = scenarios[
        (scenarios["line_id"] == top["line_id"])
        & (scenarios["scenario_reduction_pct"] == 10)
    ].iloc[0]
    category_total = down["downtime_minutes"].sum()
    unplanned_total = auto["downtime_unplanned_minutes"].sum()
    positive_records = (auto["bottleneck_gap_seconds"] > 0).mean() * 100
    mean_actual_gap = auto["actual_cycle_gap_seconds"].mean()
    num_lines_meeting_mean_takt = int((lines["mean_bottleneck_gap_seconds"] <= 0).sum())

    text = f"""# MFG-005 Automotive Line Analysis

> All observations below describe a synthetic dataset, not measured factory results.

## Scope and data quality

- Source records: {quality['source_rows']:,} rows and {quality['source_columns']:,} columns.
- Automotive subset: {quality['automotive_rows']:,} source records across {quality['automotive_lines']} production lines.
- Exact duplicate rows in source: {quality['exact_duplicate_rows']:,}.
- Repeated shift_id + line_id keys: {quality['duplicate_shift_line_keys']:,} extra source records across {quality['duplicate_shift_line_key_groups']:,} key groups in the full dataset; within automotive, {quality['automotive_duplicate_shift_line_records']:,} extra records across {quality['automotive_duplicate_shift_line_key_groups']:,} key groups.
- Missing source cells: {quality['source_missing_cells']:,}; coolant consumption is sparsely populated and is excluded from conclusions.
- Granularity: records are shift/line observations with production-order context. Automotive contains {quality['automotive_unique_shift_line_keys']:,} unique shift_id + line_id keys, so the {quality['automotive_rows']:,} rows are not all unique shifts. Repeated keys have different production orders and are retained rather than dropped.
- Date parsing: {quality['invalid_shift_dates']} invalid dates; valid automotive records span {auto['shift_date'].min():%Y-%m-%d} through {auto['shift_date'].max():%Y-%m-%d}.

## Findings

- Mean bottleneck cycle time is {auto['bottleneck_cycle_time_seconds'].mean():.2f} s versus mean takt of {auto['takt_time_seconds'].mean():.2f} s (mean gap {auto['bottleneck_gap_seconds'].mean():.2f} s).
- {positive_records:.1f}% of automotive records have a positive bottleneck gap. {num_lines_meeting_mean_takt} of {len(lines)} lines have a non-positive mean gap.
- Mean actual cycle time is {auto['actual_cycle_time_avg_seconds'].mean():.2f} s; its mean gap to takt is {mean_actual_gap:.2f} s.
- Automotive records report {auto['actual_production_quantity'].sum():,.0f} actual units, {auto['good_units_produced'].sum():,.0f} good units, and {auto['defective_units_produced'].sum():,.0f} defective units. Mean reported throughput is {auto['throughput_rate_actual_uph'].mean():.2f} units/hour.
- Mean record-level average WIP queue is {auto['wip_queue_avg_units'].mean():.2f} units (mean reported maximum WIP {auto['wip_queue_max_units'].mean():.2f} units).
- Highest mean line severity: {top['line_id']} at {top['mean_bottleneck_gap_seconds']:.2f} s above takt; bottleneck station most frequently recorded for this line: {top['dominant_bottleneck_station']}.
- That line is over takt on {top['records_over_takt_pct']:.1f}% of its records; its mean OEE is {top['mean_oee']:.1%}.
- Mean OEE is {auto['oee_overall'].mean():.1%}. The largest average component shortfall is {lead_factor['factor'].lower()} (mean factor {lead_factor['mean_factor']:.1%}, {lead_factor['mean_loss_to_100_pct']:.1f} percentage-point gap to 100%).
- The dataset's recorded primary OEE loss driver is {lead_driver['recorded_primary_loss_driver']} ({int(lead_driver['observation_count'])} records, {lead_driver['share_of_observations_pct']:.1f}%).
- Among the supplied downtime categories, {lead_down['category'].lower()} is largest at {lead_down['downtime_minutes']:,.1f} minutes ({lead_down['share_of_categorized_downtime_pct']:.1f}% of categorized downtime).
- The six supplied category fields sum to {category_total:,.1f} minutes, compared with {unplanned_total:,.1f} minutes of unplanned downtime; the {unplanned_total - category_total:,.1f}-minute mismatch means the category Pareto is not a full reconciliation.
- Average record-level unplanned downtime is {auto['downtime_unplanned_minutes'].mean():.1f} min, planned downtime {auto['downtime_planned_minutes'].mean():.1f} min, downtime events {auto['downtime_event_count'].mean():.2f}, MTBF {auto['mtbf_minutes'].mean():.1f} min, and MTTR {auto['mttr_minutes'].mean():.1f} min.

## Root-cause limits and Lean priorities

- Prioritize a time study and work-balance review at the highest-gap lines, especially the line/station named above. Verify cycle elements, starvation/blocking, and standard work before assigning a cause.
- Recorded primary OEE-loss labels point to performance loss, while supplied downtime category fields identify mechanical downtime as the largest categorized source. Neither classification proves why losses occurred.
- A focused TPM/maintenance review of mechanical events is a reasonable investigation. Validate event histories, failure modes, and maintenance records before prescribing a specific repair or claiming equipment condition caused the losses.
- Mean changeover time is {auto['changeover_time_minutes'].mean():.1f} min per record (mean changeover count {auto['changeover_count'].mean():.2f}). Equipment age/condition, operator utilization, and quality fields are available for descriptive comparison. Correlation or group differences do not establish causation; the data do not support a complete 5-Why chain.

## Hypothetical improvement simulation

- Scenario: reduce each recorded bottleneck cycle time by 10%. For {top['line_id']}, mean bottleneck time would be {ten_pct['mean_scenario_bottleneck_seconds']:.2f} s and mean gap {ten_pct['mean_scenario_gap_seconds']:.2f} s; {ten_pct['records_meeting_scenario_takt_pct']:.1f}% of its records would meet takt under this cycle-time-only scenario.
- Theoretical bottleneck-limited capacity rate for that line increases by {ten_pct['hypothetical_capacity_delta_pct']:.1f}% ({ten_pct['baseline_theoretical_capacity_uph']:.2f} to {ten_pct['scenario_theoretical_capacity_uph']:.2f} units/hour). This is a rate ceiling, not observed throughput or a production forecast.
- Absolute capacity units are not projected: repeated line-shift records have separate production orders, and the data do not clarify whether available-time values can be added. The scenario also holds downtime, quality, staffing, demand, material availability, and downstream constraints constant.

## Deliverables

- `outputs/automotive_analysis.xlsx`: data-quality checks, source data types, line ranking, bottleneck-station and downtime Pareto tables, OEE, exploratory comparisons, trends, and scenarios.
- `outputs/automotive_clean.csv`: automotive records with derived takt-gap fields.
- `outputs/powerbi_dashboard_data.csv` and `PowerBI_Measures.dax`: dashboard-ready data and suggested Power BI measures.
- `outputs/analysis_summary.json`: machine-readable headline results.
"""
    (OUTPUT / "PROJECT_FINDINGS.md").write_text(text, encoding="utf-8")


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="17365D")
    font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 32
        for cells in sheet.columns:
            letter = get_column_letter(cells[0].column)
            sample = [str(cell.value) for cell in cells[:150] if cell.value is not None]
            sheet.column_dimensions[letter].width = min(
                max(max((len(value) for value in sample), default=10) + 2, 12), 36
            )
        if sheet.max_row > 1 and sheet.max_column:
            table = Table(
                displayName=f"Table_{sheet.title.replace(' ', '')[:20]}",
                ref=sheet.dimensions,
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)
    workbook.save(path)


def main() -> None:
    OUTPUT.mkdir(exist_ok=True)
    source, auto, quality = load_dataset()
    lines = line_summary(auto)
    stations = station_summary(auto)
    down = downtime_pareto(auto)
    factors, drivers = oee_tables(auto)
    scenarios = simulation_table(auto)
    correlations, comparisons = exploratory_tables(auto)
    trend = daily_trend(auto)

    shift_columns = [
        "shift_id", "production_order_id", "shift_date", "shift_number", "line_id",
        "plant_id", "work_center_id", "product_family", "bottleneck_station_id",
        "takt_time_seconds", "actual_cycle_time_avg_seconds", "bottleneck_cycle_time_seconds",
        "bottleneck_gap_seconds", "actual_cycle_gap_seconds", "meets_takt",
        "bottleneck_utilisation_pct", "planned_production_quantity",
        "actual_production_quantity", "good_units_produced", "defective_units_produced",
        "defect_rate_pct", "throughput_rate_actual_uph", "wip_queue_avg_units",
        "wip_queue_max_units", "oee_availability", "oee_performance", "oee_quality",
        "oee_overall", "oee_loss_primary_driver", "downtime_unplanned_minutes",
        "downtime_planned_minutes", "downtime_event_count", "mtbf_minutes", "mttr_minutes",
        "changeover_time_minutes", "changeover_count", "equipment_failure_flag",
        "equipment_age_years", "equipment_condition_score", "operator_utilisation_pct",
        "primary_defect_type",
    ]
    detail = auto[[column for column in shift_columns if column in auto.columns]]
    dashboard_columns = [
        "shift_id", "production_order_id", "shift_date", "line_id",
        "bottleneck_station_id", "takt_time_seconds", "actual_cycle_time_avg_seconds",
        "bottleneck_cycle_time_seconds", "bottleneck_gap_seconds",
        "actual_production_quantity", "good_units_produced", "defective_units_produced",
        "throughput_rate_actual_uph", "wip_queue_avg_units", "oee_availability",
        "oee_performance", "oee_quality", "oee_overall", "downtime_unplanned_minutes",
        "downtime_planned_minutes", "downtime_event_count", "mtbf_minutes", "mttr_minutes",
        "changeover_time_minutes", "oee_loss_primary_driver",
        *[field for field in DOWNTIME_FIELDS if field in auto.columns],
    ]
    dashboard = auto[[column for column in dashboard_columns if column in auto.columns]]
    auto.to_csv(OUTPUT / "automotive_clean.csv", index=False, date_format="%Y-%m-%d")
    dashboard.to_csv(
        OUTPUT / "powerbi_dashboard_data.csv", index=False, date_format="%Y-%m-%d"
    )

    top = lines.iloc[0]
    summary: dict[str, object] = {
        "synthetic_data": True,
        "automotive_source_records": int(len(auto)),
        "automotive_lines": int(auto["line_id"].nunique()),
        "automotive_unique_shift_line_keys": quality["automotive_unique_shift_line_keys"],
        "automotive_repeated_shift_line_key_groups": quality[
            "automotive_duplicate_shift_line_key_groups"
        ],
        "mean_oee_pct": float(auto["oee_overall"].mean() * 100),
        "mean_takt_seconds": float(auto["takt_time_seconds"].mean()),
        "mean_actual_cycle_seconds": float(auto["actual_cycle_time_avg_seconds"].mean()),
        "mean_bottleneck_cycle_seconds": float(auto["bottleneck_cycle_time_seconds"].mean()),
        "mean_bottleneck_gap_seconds": float(auto["bottleneck_gap_seconds"].mean()),
        "records_over_takt_pct": float((auto["bottleneck_gap_seconds"] > 0).mean() * 100),
        "mean_actual_throughput_uph": float(auto["throughput_rate_actual_uph"].mean()),
        "total_actual_units": int(auto["actual_production_quantity"].sum()),
        "total_good_units": int(auto["good_units_produced"].sum()),
        "total_defective_units": int(auto["defective_units_produced"].sum()),
        "mean_wip_queue_units": float(auto["wip_queue_avg_units"].mean()),
        "mean_unplanned_downtime_minutes": float(auto["downtime_unplanned_minutes"].mean()),
        "total_unplanned_downtime_minutes": float(auto["downtime_unplanned_minutes"].sum()),
        "total_planned_downtime_minutes": float(auto["downtime_planned_minutes"].sum()),
        "mean_mtbf_minutes": float(auto["mtbf_minutes"].mean()),
        "mean_mttr_minutes": float(auto["mttr_minutes"].mean()),
        "highest_mean_gap_line": str(top["line_id"]),
        "highest_mean_gap_station": str(top["dominant_bottleneck_station"]),
        "highest_mean_gap_seconds": float(top["mean_bottleneck_gap_seconds"]),
        "largest_recorded_oee_loss_driver": str(drivers.iloc[0]["recorded_primary_loss_driver"]),
        "largest_downtime_category": str(down.iloc[0]["category"]),
        "source_data_quality": quality,
    }
    (OUTPUT / "analysis_summary.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )
    write_findings(auto, quality, lines, down, factors, drivers, scenarios)

    read_me = pd.DataFrame(
        {
            "Topic": [
                "Data scope", "Observation granularity", "Shift detail", "Bottleneck gap",
                "Severity ranking", "OEE", "Downtime Pareto", "Simulation",
                "Root-cause interpretation",
            ],
            "Definition / limitation": [
                "Synthetic source filtered to industry_sector = automotive.",
                "Source rows carry production-order context; repeated shift_id + line_id keys are retained because production orders differ.",
                "Record-level values; line means are unweighted means, so repeated line-shift keys can weight a shift more than once.",
                "Bottleneck cycle time minus takt; positive means recorded bottleneck time exceeds takt.",
                "Lines ranked by mean bottleneck gap; Pareto shares use summed positive record-level gaps.",
                "Source OEE factor values are ratios from 0 to 1; component shortfalls compare means to 100%.",
                "Aggregates six supplied category fields; the category sum may not equal unplanned downtime.",
                "5%, 10%, and 15% cycle-time reductions are hypothetical capacity-rate ceilings; no absolute units are projected.",
                "Descriptive records and correlations do not prove cause; validate suspected causes on the production floor.",
            ],
        }
    )
    data_quality = [
        {"check": "Source rows", "value": quality["source_rows"]},
        {"check": "Source columns", "value": quality["source_columns"]},
        {"check": "Exact duplicate rows", "value": quality["exact_duplicate_rows"]},
        {"check": "Extra duplicate shift_id + line_id source records", "value": quality["duplicate_shift_line_keys"]},
        {"check": "Source repeated shift_id + line_id key groups", "value": quality["duplicate_shift_line_key_groups"]},
        {"check": "Extra automotive shift_id + line_id records", "value": quality["automotive_duplicate_shift_line_records"]},
        {"check": "Automotive repeated shift_id + line_id key groups", "value": quality["automotive_duplicate_shift_line_key_groups"]},
        {"check": "Automotive source records", "value": quality["automotive_rows"]},
        {"check": "Automotive lines", "value": quality["automotive_lines"]},
        {"check": "Invalid parsed shift dates", "value": quality["invalid_shift_dates"]},
    ]
    data_quality.extend(
        {"check": f"Source dtype: {dtype}", "value": count}
        for dtype, count in quality["source_dtype_counts"].items()
    )
    data_quality.extend(
        {"check": f"Missing: {column}", "value": count}
        for column, count in quality["source_missing_by_column"].items()
    )
    with pd.ExcelWriter(OUTPUT / "automotive_analysis.xlsx", engine="openpyxl") as writer:
        read_me.to_excel(writer, sheet_name="Read Me", index=False)
        pd.DataFrame(data_quality).to_excel(writer, sheet_name="Data Quality", index=False)
        pd.DataFrame(
            {
                "column": source.columns,
                "inferred_pandas_dtype": source.dtypes.astype(str).to_numpy(),
                "missing_cells": source.isna().sum().to_numpy(),
            }
        ).to_excel(writer, sheet_name="Source Data Types", index=False)
        lines.to_excel(writer, sheet_name="Line Summary", index=False)
        stations.to_excel(writer, sheet_name="Bottleneck Stations", index=False)
        down.to_excel(writer, sheet_name="Downtime Pareto", index=False)
        factors.to_excel(writer, sheet_name="OEE Factors", index=False)
        drivers.to_excel(writer, sheet_name="OEE Drivers", index=False)
        comparisons.to_excel(writer, sheet_name="Loss Comparisons", index=False)
        correlations.to_excel(writer, sheet_name="Spearman Correlations", index=False)
        trend.to_excel(writer, sheet_name="Daily Trend", index=False)
        scenarios.to_excel(writer, sheet_name="SIM Scenario", index=False)
        detail.to_excel(writer, sheet_name="Shift Detail", index=False)
    style_workbook(OUTPUT / "automotive_analysis.xlsx")

    print(f"Source rows: {len(source):,}; automotive records: {len(auto):,}")
    print(f"Automotive lines: {auto['line_id'].nunique()}")
    print(f"Mean OEE: {summary['mean_oee_pct']:.1f}%")
    print(
        f"Mean bottleneck gap: {summary['mean_bottleneck_gap_seconds']:.2f} s; "
        f"records over takt: {summary['records_over_takt_pct']:.1f}%"
    )
    print(f"Highest-severity line: {summary['highest_mean_gap_line']}")
    print(f"Workbook: {OUTPUT / 'automotive_analysis.xlsx'}")


if __name__ == "__main__":
    main()
